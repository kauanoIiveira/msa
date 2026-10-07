from pathlib import Path
import csv, hashlib, json, shutil, unicodedata, zipfile, re, math, subprocess
from collections import Counter
from datetime import datetime, date
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(r'C:\Users\Kauan\Documents\ChatGPT\Desafio de Ideias')
SOURCE = Path(r'C:\Users\Kauan\Desktop\Grupo Amarelo - MSA Brasil')
DESKTOP = Path(r'C:\Users\Kauan\Desktop\msa-master')
OUT = ROOT / 'docs/continuidade/2026-10-06'
EVID = OUT / 'evidencias'
EVID.mkdir(parents=True, exist_ok=True)

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def write(name, obj):
    (EVID / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

prompt = Path(r'C:\Users\Kauan\.codex\attachments\989777ab-36cb-43d2-9933-7404448bbfb0\Texto colado.txt')
shutil.copyfile(prompt, OUT / 'PROMPT_ORIGINAL_INTEGRAL.txt')
assert digest(prompt) == digest(OUT / 'PROMPT_ORIGINAL_INTEGRAL.txt')

inventory = []
for path in sorted(SOURCE.rglob('*')):
    if not path.is_file():
        continue
    item = dict(path=str(path), relative=str(path.relative_to(SOURCE)), size=path.stat().st_size, sha256=digest(path), type=path.suffix.lower())
    inventory.append(item)
write('inventario_local.json', inventory)

prior = json.loads((ROOT / 'analise/evidencias/fontes_sha256.json').read_text(encoding='utf-8'))
matches = []
for item in prior:
    same = [p for p in inventory if Path(p['path']).name == Path(item['path']).name]
    matches.append(dict(previous=item, current=same, identical=bool(same) and all(p['sha256'] == item['sha256'] for p in same)))
write('conferencia_extracoes_anteriores.json', matches)

desktop_files = {str(p.relative_to(DESKTOP)).replace('\\','/'): p for p in DESKTOP.rglob('*') if p.is_file()}
comparison = dict(total_desktop=len(desktop_files), identical=[], different=[], desktop_only=[])
for relative,path in desktop_files.items():
    other=ROOT/relative
    record=dict(relative=relative,desktop_sha256=digest(path),current_sha256=digest(other) if other.is_file() else None)
    comparison['desktop_only' if not other.is_file() else 'identical' if record['desktop_sha256']==record['current_sha256'] else 'different'].append(record)
write('comparacao_checkouts.json',comparison)

ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
docs=[]
for path in sorted(SOURCE.rglob('*.docx')) + sorted((ROOT/'docs').glob('*.docx')):
    sections=[]
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith('word/') and name.endswith('.xml') and any(s in name for s in ['document.xml','header','footer','comments','footnotes','endnotes']):
                tree=ET.fromstring(z.read(name))
                paragraphs=[]
                for p in tree.findall('.//w:p',ns):
                    runs=[]
                    for e in p.iter():
                        if e.tag == '{'+ns['w']+'}t': runs.append(e.text or '')
                        elif e.tag == '{'+ns['w']+'}tab': runs.append('\t')
                        elif e.tag == '{'+ns['w']+'}br': runs.append('\n')
                    paragraphs.append(''.join(runs))
                sections.append(dict(part=name,text='\n'.join(paragraphs)))
        media=[n for n in z.namelist() if n.startswith('word/media/')]
    text='\n\n'.join('## '+s['part']+'\n'+s['text'] for s in sections)
    name=f'docx_{len(docs)+1:02}.txt'
    (EVID/name).write_text(text,encoding='utf-8')
    docs.append(dict(path=str(path),extraction=name,chars=len(text),media=media))
write('indice_docx.json',docs)

pdfs=[]
for path in sorted(SOURCE.rglob('*.pdf')):
    reader=PdfReader(path)
    texts=[]
    pages=[]
    for i,page in enumerate(reader.pages):
        content=page.extract_text() or ''
        texts.append(f'--- PAGINA {i+1} ---\n{content}')
        pages.append(dict(page=i+1,chars=len(content),images=len(page.images)))
    name=path.stem+'.txt'
    (EVID/name).write_text('\n\n'.join(texts),encoding='utf-8')
    pdfs.append(dict(path=str(path),extraction=name,pages=pages))
write('indice_pdfs.json',pdfs)

workbooks=[]
for path in sorted(SOURCE.rglob('*.xlsx')):
    values=load_workbook(path,data_only=True,read_only=False)
    formulas=load_workbook(path,data_only=False,read_only=False)
    sheets=[]
    for ws in formulas.worksheets:
        cells=[]
        for row in ws:
            for c in row:
                if c.value is not None or c.comment:
                    cells.append(dict(address=c.coordinate,value=c.value,type=c.data_type,cached=values[ws.title][c.coordinate].value,number_format=c.number_format,comment=c.comment.text if c.comment else None))
        name='planilha_'+ws.title.replace(' ','_')+'.json'
        write(name,cells)
        sheets.append(dict(name=ws.title,max_row=ws.max_row,max_column=ws.max_column,nonempty=len(cells),formulas=sum(c['type']=='f' for c in cells),merged=[str(x) for x in ws.merged_cells],extraction=name))
    workbooks.append(dict(path=str(path),sha256=digest(path),sheets=sheets,defined_names=str(formulas.defined_names),external_links=len(formulas._external_links)))
write('indice_planilha.json',workbooks)

csvs=[]
for path in sorted(SOURCE.rglob('*.csv')):
    text=path.read_text(encoding='utf-8-sig')
    dialect=csv.Sniffer().sniff(text[:5000])
    rows=list(csv.DictReader(text.splitlines(),dialect=dialect))
    csvs.append(dict(path=str(path),rows=len(rows),fields=list(rows[0]) if rows else [],sha256=digest(path),matches_checkout=any(digest(p)==digest(path) for p in (ROOT/'entrega').glob(path.name))))
write('indice_csv.json',csvs)
if (EVID/'arquivos_drive_materializados.json').exists():
    manifest=json.loads((EVID/'arquivos_drive_materializados.json').read_text(encoding='utf-8'))
    remote=[]
    for item in manifest:
        path=Path(item['path'])
        sha=digest(path)
        matching=[p['path'] for p in inventory if p['sha256']==sha]
        remote.append(dict(**item,size=path.stat().st_size,sha256=sha,identical_local_files=matching))
    write('conferencia_drive_local.json',remote)
    print(json.dumps(dict(drive_files=len(remote),drive_identical_local=sum(bool(x['identical_local_files']) for x in remote),drive_differences=[x for x in remote if not x['identical_local_files']]),ensure_ascii=False))

for item in comparison['different']:
    a=desktop_files[item['relative']].read_bytes().replace(b'\r\n',b'\n')
    b=(ROOT/item['relative']).read_bytes().replace(b'\r\n',b'\n')
    item['same_after_line_ending_normalization']=a==b
comparison['line_endings_only']=sum(x.get('same_after_line_ending_normalization',False) for x in comparison['different'])
comparison['content_differences']=[x['relative'] for x in comparison['different'] if not x['same_after_line_ending_normalization']]
comparison['current_git_only']=[x for x in subprocess.check_output(['git','ls-files'],cwd=ROOT,text=True).splitlines() if x not in desktop_files]
write('comparacao_checkouts.json',comparison)
ws=load_workbook(next(SOURCE.rglob('*.xlsx')),data_only=False)['Selo ']
cols=[c.column for c in ws[9] if c.value is not None and c.column>=6]
counts=Counter()
for row in range(17,34):
    for col in cols:
        c=ws.cell(row,col)
        counts['missing' if c.value is None else 'number' if isinstance(c.value,(int,float)) else 'text']+=1
summary=dict(parameters=len(cols),collections=17,fields=counts,dates=[str(ws.cell(r,1).value) for r in range(17,34)],limits={ws.cell(9,c).coordinate:dict(name=ws.cell(9,c).value,lower=ws.cell(10,c).value,upper=ws.cell(11,c).value,unit=ws.cell(12,c).value) for c in cols},formula_samples={coord:ws[coord].value for coord in ['BF69','BF71','BF72','BF73','BF74','BH10','BH11','U6','G93']})
write('resumo_planilha_conferido.json',summary)
print(json.dumps(dict(sources=len(inventory),prior_identical=sum(x['identical'] for x in matches),desktop_identical=len(comparison['identical']),desktop_line_endings_only=comparison['line_endings_only'],desktop_content_differences=comparison['content_differences'],current_git_only=comparison['current_git_only'],spreadsheet=summary),ensure_ascii=False,default=str))
