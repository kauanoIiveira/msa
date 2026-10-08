from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/BRIEFING_PITCH_MSA_2026-10-07.md'
OUTPUT = ROOT / 'docs/Briefing_Pitch_MSA_2026-10-07.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin, section.bottom_margin = Inches(.7), Inches(.65)
section.left_margin, section.right_margin = Inches(.8), Inches(.8)
section.header_distance, section.footer_distance = Inches(.3), Inches(.3)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Header', 'Footer']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(9)
for style_name, size in [('Heading 1', 18), ('Heading 2', 12)]:
    style = doc.styles[style_name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.text = 'MSA Brasil    Equipe Amarela    Briefing para o pitch'
header.runs[0].font.size = Pt(9)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('07/10/2026    ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(9)

def inline(paragraph, text):
    for token in re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))', text):
        if token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif re.fullmatch(r'\[[^\]]+\]\(https?://[^)]+\)', token):
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token).groups()
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            properties = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '005A70'); properties.append(color)
            run.append(properties)
            node = OxmlElement('w:t'); node.text = label; run.append(node)
            hyperlink.append(run); paragraph._p.append(hyperlink)
        else:
            paragraph.add_run(token.replace('`', ''))

def table(rows):
    widths = {
        3: [2.65, 2.55, 1.7],
        4: [1.75, 1.85, 1.75, 1.55],
        2: [2.25, 4.65],
    }[len(rows[0])]
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element = OxmlElement('w:' + edge)
        for key, val in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            element.set(qn('w:' + key), val)
        borders.append(element)
    props.append(borders)
    for i, values in enumerate(rows):
        row = t.add_row()
        prevent_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(prevent_split)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_props = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                edge_element = OxmlElement('w:' + edge)
                edge_element.set(qn('w:w'), '85'); edge_element.set(qn('w:type'), 'dxa')
                margins.append(edge_element)
            cell_props.append(margins)
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E4EBEE' if i == 0 else ('F5F7F8' if i % 2 == 0 else 'FFFFFF'))
            cell_props.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.02
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(10)
                if i == 0:
                    run.bold = True
            if re.match(r'^R\$|^\*\*R\$', value):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
page_pending = False
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1; continue
    if line == '<!-- page -->':
        page_pending = True; i += 1; continue
    if line.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            values = [part.strip() for part in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?', value) for value in values):
                rows.append(values)
            i += 1
        table(rows); continue
    if line.startswith('# '):
        p = doc.add_paragraph(style='Title'); inline(p, line[2:])
    elif line.startswith('## '):
        p = doc.add_paragraph(style='Heading 1'); inline(p, line[3:])
    elif line.startswith('### '):
        p = doc.add_paragraph(style='Heading 2'); inline(p, line[4:])
    elif line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet'); inline(p, line[2:])
    elif re.match(r'^\d+\. ', line):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.16)
        p.paragraph_format.first_line_indent = Inches(-.16)
        inline(p, line)
    else:
        p = doc.add_paragraph(); inline(p, line)
    if page_pending:
        p.paragraph_format.page_break_before = True
        page_pending = False
    i += 1

doc.core_properties.title = 'Conteudo para o pitch do sistema MSA'
doc.core_properties.author = 'Equipe Amarela'
doc.core_properties.subject = 'Briefing com cenario simulado e viabilidade de referencia'
doc.save(OUTPUT)
print(str(OUTPUT))
