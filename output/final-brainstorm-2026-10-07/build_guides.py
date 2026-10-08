from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]

def build(source, output, subject):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Header', 'Footer']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        for border in list(style.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.04
    doc.styles['Title'].font.size = Pt(24)
    for name, size in [('Heading 1', 16), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = 'MSA Brasil    Equipe Amarela    ' + subject
    header.runs[0].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('07/10/2026    ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    pending = False
    for line in (ROOT / source).read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if not line:
            continue
        if line == '<!-- page -->':
            pending = source.startswith('docs/ROTEIRO_')
            continue
        style = 'Normal'
        for prefix, selected in [('### ', 'Heading 2'), ('## ', 'Heading 1'), ('# ', 'Title')]:
            if line.startswith(prefix):
                line, style = line[len(prefix):], selected
                break
        paragraph = doc.add_paragraph(style=style)
        if pending:
            paragraph.paragraph_format.page_break_before = True
            pending = False
        for part in re.split(r'(\*\*.*?\*\*)', line):
            run = paragraph.add_run(part[2:-2] if part.startswith('**') else part)
            run.bold = part.startswith('**')
    doc.core_properties.title = doc.paragraphs[0].text
    doc.core_properties.author = 'Equipe Amarela'
    doc.core_properties.subject = subject
    doc.save(ROOT / output)
    print(ROOT / output)

build('docs/GUIA_FUNCOES_MSA_2026-10-07.md', 'docs/Guia_Funcoes_MSA_2026-10-07.docx', 'Guia de uso')
build('docs/ROTEIRO_PAGINAS_MSA_2026-10-07.md', 'docs/Roteiro_Paginas_MSA_2026-10-07.docx', 'Roteiro por pagina')
