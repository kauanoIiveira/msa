from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'Roteiro_Entrevista_MSA_Equipe_Amarela.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for style_name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[style_name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for border in style.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Equipe Amarela | MSA | ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    p(text, 'Heading 2')

def question(title, text):
    paragraph = p('')
    paragraph.add_run(title + ' ').bold = True
    paragraph.add_run(text)

p('Roteiro da entrevista com a MSA', 'Title')
p('Equipe Amarela | Preparação em 05 de outubro de 2026')
p('Objetivo: usar 20 minutos com uma pessoa da empresa para entender a rotina, confirmar as regras e combinar o primeiro piloto. O MVP ajuda a conversar; nossas hipóteses ainda dependem da empresa.')
h('O que entendemos do problema')
p('Hoje, parte dos parâmetros sai de fotos da IHM e chega ao Excel por transcrição. A equipe também precisa reunir produção, paradas, motivos, refugos e perdas. A Engenharia consulta esses registros para conferir limites, acompanhar o processo e decidir sobre a produção.')
p('Nossa proposta concentra o registro e a consulta em um fluxo rastreável. Queremos reduzir a repetição de tarefas sem aumentar o trabalho de quem está na máquina. Ainda precisamos ouvir como os turnos funcionam e quais informações a equipe considera confiáveis.')
h('Uma abertura que podemos usar')
p('“Pelo material, entendemos que vocês precisam juntar os parâmetros da máquina com o que aconteceu na produção. Antes de mostrar o sistema, vocês podem nos contar como fazem isso em um turno e onde encontram mais dificuldade? Trouxemos um fluxo para vocês avaliarem, e queremos ajustá-lo à rotina de vocês.”')
h('Divisão entre os cinco integrantes')
p('Uma pessoa conduz e controla o tempo. Uma pergunta sobre operação; outra, sobre regras e Engenharia; outra, sobre dispositivos e TI. A quinta registra respostas, pendências e responsáveis. Os especialistas entram quando o condutor passa a palavra. Não fazer uma rodada em que todos repetem perguntas.')
h('Ritmo dos 20 minutos')
p('0 a 2 min: apresentação e confirmação do objetivo. 2 a 7 min: rotina e principal dificuldade. 7 a 11 min: quantidades e regras. 11 a 14 min: dispositivos e rede. 14 a 17 min: prioridades e piloto. 17 a 20 min: confirmar entendimento e próximos passos.')
h('Demonstração apenas se ajudar a conversa')
p('Se a pessoa concordar, reservar até 90 segundos dentro do bloco de prioridades: selecionar máquina e produto, mostrar um desvio e a fila da Engenharia. Dizer que os números são fictícios e perguntar se esse fluxo faz sentido. Não tentar apresentar todas as telas durante a entrevista.')
h('O que podemos afirmar')
p('O MVP cobre os entregáveis obrigatórios por cadastro, apontamento manual, armazenamento no Firebase, histórico e dashboard. Os dados fictícios permitem avaliar o fluxo; não comprovam desempenho da fábrica. Não temos conexão com as máquinas, OCR ou coleta automática. Uma aprovação no sistema não aciona nem libera fisicamente o equipamento.')

doc.add_page_break()
p('Perguntas prioritárias', 'Heading 1')
p('Sete perguntas guiam a conversa. Usar as perguntas técnicas da próxima página apenas como apoio. Se a pessoa já responder um ponto, não perguntar de novo. Não exigir números que ela não tenha em mãos.')
question('Rotina atual', 'Quem coleta, transcreve e confere os dados? Em quais momentos do turno? Podem nos mostrar o caminho de uma foto até a decisão da Engenharia?')
question('Maior dificuldade', 'Onde há mais retrabalho ou espera hoje? Qual informação costuma faltar quando vocês precisam decidir?')
question('Quantidades e paradas', 'A produção registrada é bruta ou boa? Como vocês registram refugos, perdas e o início e fim de uma parada? Pedir um exemplo e aprofundar só o ponto que causar dúvida.')
question('Regras e aprovação', 'Os limites mudam por produto ou receita? Quem valida esses limites e o que precisa acontecer quando uma leitura fica fora da faixa?')
question('Tablets e acesso', 'A empresa já disponibiliza tablet, celular ou computador junto às máquinas? Como é o acesso à rede no local? Perguntar sobre câmera apenas se OCR for prioridade.')
question('Diferenciais desejados', 'Dos opcionais, quais dois ou três resolveriam melhor a dificuldade que vocês descreveram? Se precisar, citar OCR, coleta direta, QR, alertas e relatórios. Há algum indispensável que não está na folha?')
question('Primeiro piloto e sucesso', 'Onde começariam e quem poderia nos ajudar a validar? Qual melhoria mostraria que o piloto valeu a pena?')
h('Condução da conversa')
p('Deixar a pessoa terminar a resposta. Pedir “pode dar um exemplo?” antes de sugerir uma solução. Se ela não responder por Engenharia ou TI, registrar a pendência e pedir um contato, sem insistir. Reservar os três minutos finais mesmo que sobrem perguntas.')
h('Anotar antes de encerrar')
p('Responsável pelo piloto: ______________________    Máquina e produto: ______________________')
p('Prioridades escolhidas: ______________________________________________________________')
p('Próxima ação e responsável: ______________________________    Prazo: _____________________')

doc.add_page_break()
p('Dúvidas técnicas e próximos passos', 'Heading 1')
p('Se faltar tempo, enviar estas perguntas para a Engenharia ou TI. Não precisamos resolver tudo durante a primeira conversa, mas precisamos saber quem pode responder.')
question('Aquecimento Z1 a Z21', 'A planilha registra temperatura medida ou valor configurado? As zonas com limites 0 e 0 estão desligadas, sem referência ou possuem outro significado? As faixas variam com a receita?')
question('Contramolde', 'O tempo aparece com mínimo 80 e máximo 75 segundos. Há troca de campos, outra unidade, duas receitas ou uma interpretação diferente? Qual referência devemos usar?')
question('Vácuo e pressão', 'Para o vácuo de -600 mm/Hg, a condição aceita valores mais negativos ou menos negativos? A pressão tem apenas mínimo de 6,5 bar ou também um máximo?')
question('Médias e capacidade', 'O “measure mean” representa qual grandeza e unidade? Como a Engenharia define amostragem, média, desvio, Cp/Cpk e critério de aprovação? Há subgrupos e validação do instrumento?')
question('Integração e infraestrutura', 'Quais CLPs e IHMs estão envolvidos? A TI permite integração ou exportação de dados? O sistema deve usar banco, autenticação ou servidor da empresa? Como agir quando a rede cair?')
question('Alertas e histórico', 'Quem deve receber alertas, por qual canal e em qual horário? Por quanto tempo guardam registros e fotos? Precisam de relatório específico, trilha de auditoria ou integração com outro sistema?')
h('Cuidados que demonstram entendimento')
for text in [
    'Não inverter o limite do contramolde nem definir a direção do vácuo por conta própria. Não tratar 0 e 0 como zona desligada sem confirmação.',
    'Separar peças de kg e produção bruta de produção boa. Manter motivo, período, autor e versão do limite ligados ao registro.',
    'Tratar campo vazio como informação ausente. A planilha contém números como texto; a entrada deve validar vírgula e ponto antes de calcular.',
    'A auditoria reproduziu os cálculos com desvio populacional. Isso não valida estabilidade, normalidade ou capacidade industrial. Confirmar o método com a Engenharia.'
]:
    p(text, 'List Bullet')
h('Como fechar a conversa')
p('“Para confirmar se entendemos: o primeiro piloto será em ___, a principal dificuldade é ___ e vocês priorizam ___. Podemos combinar quem valida os limites e um exemplo de turno para testar o fluxo? Vamos registrar as pendências e retornar com os ajustes, sem pedir que o operador faça trabalho duplicado.”')
p('Base deste roteiro: folha do Desafio de Ideias MSA, planilha T20A03(EN)5 e auditoria local dos materiais. Os pontos apresentados como perguntas ainda dependem de confirmação da empresa.')
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
