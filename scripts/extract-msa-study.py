"""Read-only extraction of the supplied capability workbook for the MSA reference viewer."""
import argparse
import hashlib
import json
from datetime import date, datetime
from pathlib import Path
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument('source', type=Path)
parser.add_argument('--output', type=Path, required=True)
args = parser.parse_args()
original = args.source.read_bytes()
digest = hashlib.sha256(original).hexdigest()
formula_book = openpyxl.load_workbook(args.source, read_only=True, data_only=False)
value_book = openpyxl.load_workbook(args.source, read_only=True, data_only=True)
sheet = formula_book['Selo ']
values = value_book['Selo ']
formula_rows = list(sheet.iter_rows(values_only=True))
value_rows = list(values.iter_rows(values_only=True))
formula_at = lambda row, col: formula_rows[row - 1][col - 1]
value_at = lambda row, col: value_rows[row - 1][col - 1]

def serialize(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value

def normalized_date(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()[:10]
    if isinstance(value, str):
        for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return None

parameters = []
for col in range(6, sheet.max_column + 1):
    if not formula_at(9, col) or not formula_at(12, col):
        continue
    letter = openpyxl.utils.get_column_letter(col)
    samples = []
    for row in range(17, 68):
        day = value_at(row, 1)
        if day is None:
            continue
        raw = value_at(row, col)
        samples.append({'row': row, 'date': normalized_date(day), 'rawDate': serialize(day), 'cell': f'{letter}{row}', 'raw': serialize(raw)})
    parameters.append({'code': 'MSA_' + letter, 'column': letter,
        'name': str(formula_at(9, col)), 'unit': str(formula_at(12, col)),
        'limits': {'lower': serialize(value_at(10, col)), 'upper': serialize(value_at(11, col))},
        'samples': samples,
        'excelSummary': [{'cell': f'{letter}{row}', 'formula': serialize(formula_at(row, col)),
                          'cached': serialize(value_at(row, col))} for row in range(69, 75)]})
assert len(parameters) == 41, f'Unexpected source structure: {len(parameters)} parameters'
assert all(len(p['samples']) == 17 for p in parameters), 'Unexpected observation rows'
result = {'source': {'name': args.source.name, 'sheet': sheet.title, 'sha256': digest,
    'observationRows': 17, 'timestampPrecision': 'date', 'chronologyConfirmed': False,
    'sheets': [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in formula_book]},
    'parameters': parameters}
args.output.parent.mkdir(parents=True, exist_ok=True)
args.output.write_text('export const capabilityStudy = ' + json.dumps(result, ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
formula_book.close()
value_book.close()
assert hashlib.sha256(args.source.read_bytes()).hexdigest() == digest, 'Original workbook changed'
print(json.dumps({'sha256': digest, 'parameters': len(parameters), 'rows': 17,
    'numericText': sum(isinstance(s['raw'], str) for p in parameters for s in p['samples']),
    'missing': sum(s['raw'] is None for p in parameters for s in p['samples'])}))
