from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs' / 'Roteiro_Apresentacao_Funcoes_MSA.docx'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for border in style.element.xpath('./w:pPr/w:pBdr'):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.04
doc.styles['List Bullet'].font.size = Pt(11)
doc.styles['List Bullet'].paragraph_format.space_after = Pt(3)
doc.styles['Title'].font.size = Pt(22)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 2'].font.size = Pt(11.5)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(5)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
footer = s.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('Equipe Amarela | Apresentação MSA | ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def section(title, speaker, functions, speech, show, new_page=False):
    heading = p(title, 'Heading 1')
    if new_page:
        heading.paragraph_format.page_break_before = True
    p(speaker + ' | Cerca de 1 minuto incluindo os cliques')
    p('Funções disponíveis', 'Heading 2')
    for text in functions:
        p(text, 'List Bullet')
    p('Fala sugerida', 'Heading 2')
    p('“' + speech + '”')
    paragraph = p('')
    paragraph.add_run('Mostrar na tela: ').bold = True
    paragraph.add_run(show)

p('Roteiro de apresentação do sistema MSA', 'Title')
p('Sete páginas do sistema em cerca de sete minutos. Você apresenta as três primeiras; seu colega assume as demais. As listas servem de apoio, não precisam ser lidas em voz alta. Reservar o restante dos 20 minutos para ouvir a empresa.')
p('Abertura: “Os dados desta apresentação são fictícios. Queremos mostrar o fluxo e ouvir o que precisa mudar para atender à rotina de vocês.”')

section('Dashboard', 'Você', [
    'Filtrar por máquina, processo, produto e período, inclusive datas personalizadas.',
    'Consultar produção bruta, peças boas, refugo, tempo de parada e fila da Engenharia.',
    'Ver gráficos de produção diária e duração das paradas por motivo.',
    'Consultar as 21 zonas de aquecimento e os demais parâmetros do processo.',
    'Identificar desvios de parâmetros e metas fora do esperado.',
    'Abrir detalhes dos parâmetros, iniciar uma coleta e exportar coletas em CSV.',
    'Iniciar um dos 11 cenários locais e voltar aos registros do Firebase.'
], 'Aqui reunimos a visão da produção. Podemos selecionar máquina, processo, produto e período. Mostramos produção bruta, peças boas, refugos, tempo de parada e análises pendentes. Os gráficos ajudam a acompanhar a produção e os principais motivos de parada. Também aparecem os parâmetros da máquina e metas fora do esperado. O simulador permite demonstrar situações diferentes sem alterar os registros do Firebase.', 'Apontar filtros e indicadores. Mostrar um desvio; não percorrer os 11 cenários.')

section('Parâmetros', 'Você', [
    'Consultar os 41 parâmetros de referência e parâmetros adicionais cadastrados.',
    'Buscar pelo nome e consultar leitura, unidade, faixa e situação.',
    'Abrir o histórico gráfico, média, desvio populacional e estimativas de Cp e Cpk.',
    'Distinguir leitura ausente ou inválida e limite pendente, sem tratar ausência como zero.',
    'Registrar nova coleta, criar uma versão de limites quando autorizado e exportar CSV.'
], 'Nesta página consultamos os parâmetros identificados nos materiais. Cada um apresenta sua última leitura, limite e situação. Ao abrir um parâmetro, vemos o histórico, média, desvio e estimativas de Cp e Cpk quando o cálculo é aplicável. Limites ainda incertos permanecem pendentes. Também podemos registrar uma nova coleta. Precisamos confirmar com a Engenharia a interpretação das faixas e o método estatístico.', 'Abrir Medida do Passo. Apontar histórico e média; evitar explicar fórmulas durante a fala.')

section('Apontamentos', 'Você', [
    'Produção: registrar quantidade, início, fim e base bruta ou boa.',
    'Paradas: registrar início, motivo e condição planejada ou não planejada.',
    'Encerrar a parada uma vez e calcular sua duração, mantendo o motivo original.',
    'Perdas: registrar refugo, perda de material ou retrabalho, com quantidade e motivo.',
    'Separar peças de kg e manter vínculos com máquina, processo e produto.',
    'Consultar detalhes dos registros e exportar a aba selecionada em CSV.'
], 'Aqui registramos o que aconteceu na produção. Informamos quantidade, período e se são peças brutas ou boas. Em Paradas, registramos início e motivo; o encerramento permite calcular a duração. Em Perdas, distinguimos refugos, perdas de material e retrabalho. Peças e quilos ficam separados. Os registros pertencem à máquina, processo e produto selecionados, para analisar o contexto correto.', 'Alternar Produção, Paradas e Perdas. Abrir um formulário sem preenchê-lo inteiro.', new_page=True)

section('Engenharia', 'Colega', [
    'Consultar análises aguardando, em análise, aprovadas e não aprovadas.',
    'Iniciar a análise e registrar aprovação ou não aprovação com justificativa.',
    'Consultar detalhes da análise e o histórico das decisões.',
    'Consultar correções propostas, seus motivos e o registro original preservado.',
    'Decidir uma correção quando autorizado e quando o decisor não é seu autor.',
    'Exportar a aba selecionada em CSV; aplicar permissões conforme o perfil.'
], 'Esta página organiza as coletas que precisam de avaliação. A análise passa por aguardando, em análise e uma decisão, com justificativa. Também consultamos correções propostas, preservando o registro original. Uma correção exige a decisão de outra pessoa, evitando que o próprio autor aprove sua alteração. A aprovação registrada aqui não libera fisicamente a máquina; queremos confirmar como essa decisão acontece na empresa.', 'Mostrar uma análise pendente e outra decidida. Não alterar uma decisão já concluída.')
p('Atenção: solicitar uma nova correção ainda não possui um formulário completo na interface. A consulta e a decisão de propostas existentes estão disponíveis.')

section('Histórico', 'Colega', [
    'Consultar seis abas: Coletas, Produção, Paradas, Perdas, Análises e Correções.',
    'Filtrar os registros por contexto e período e abrir seus detalhes.',
    'Consultar leituras originais, situação dos registros e justificativas de decisões.',
    'Enviar uma coleta à Engenharia, definindo o escopo da análise.',
    'Encerrar uma parada ainda aberta quando autorizado.',
    'Exportar a aba selecionada em CSV para consulta fora do sistema.'
], 'Aqui consultamos coletas, produção, paradas, perdas, análises e correções do período selecionado. Podemos abrir os detalhes para entender o registro e consultar justificativas das decisões. Isso ajuda a investigar uma ocorrência sem procurar informações em lugares diferentes. A exportação gera um CSV para continuar a análise fora do sistema. Queremos saber quais consultas e relatórios vocês usam com mais frequência.', 'Abrir uma coleta e apontar as leituras. Mostrar Enviar à Engenharia sem enviar um registro real.', new_page=True)

section('Cadastros', 'Colega', [
    'Máquinas: cadastrar nome e código; processos: associar à máquina.',
    'Produtos: cadastrar e associar aos processos permitidos.',
    'Parâmetros: cadastrar nome, código e processo a ser monitorado.',
    'Limites: criar versões de unidade, natureza, status e regra aplicável.',
    'Cadastrar referências: instalar o catálogo de 41 parâmetros como rascunhos, após confirmação e definição da natureza de cada um.',
    'Motivos: cadastrar motivos de parada, refugo, perda de material e retrabalho.',
    'Metas: definir indicador, valor, mínimo ou máximo, contexto e período.',
    'Ativar ou inativar cadastros sem excluir o histórico existente.'
], 'Aqui configuramos máquinas, processos, produtos, parâmetros, motivos e metas. Esses vínculos organizam os apontamentos. Os limites possuem versões, preservando a referência usada nas coletas anteriores. As metas definem o resultado esperado para um contexto e período. Também podemos inativar cadastros. Precisamos confirmar quem poderá administrar essas informações e se os limites mudam conforme produto, receita, molde ou lote.', 'Mostrar Produtos, Parâmetros e Metas. Abrir Limites apenas para consultar o formulário.')

section('Configurações', 'Colega', [
    'Perfil: consultar e-mail e papel de acesso; editar e salvar o nome.',
    'Senha: trocar a senha após confirmar a senha atual e a nova senha.',
    'Tema: escolher claro, escuro ou acompanhar o dispositivo.',
    'VLibras: ativar ou desativar o recurso de tradução em Libras.',
    'Consulta: ativar tabelas compactas e definir período inicial de 7, 14 ou 30 dias.',
    'Restaurar preferências de consulta; manter tema e acessibilidade separados.'
], 'Aqui o usuário ajusta seu nome, troca a senha e escolhe o tema. Também pode ativar o VLibras, ajustar as tabelas e definir o período inicial de consulta. Essas preferências adaptam a consulta ao dispositivo utilizado. Para planejar o uso na operação, precisamos saber se vocês já disponibilizam tablets, celulares ou computadores próximos às máquinas.', 'Mostrar tema e período inicial. Não trocar a senha durante a entrevista.', new_page=True)

p('Funções de acesso e navegação', 'Heading 1')
for text in [
    'Conectar: entrar com e-mail e senha; mostrar ou ocultar senha; solicitar recuperação por e-mail. Não há cadastro público de usuários.',
    'Restaurar sessão e consultar permissões; Sair encerra o acesso. Dentro da simulação, Sair retorna aos registros.',
    'Menu lateral e menu móvel: trocar de página. Avatar abre Configurações; sino abre Engenharia, sem envio automático de notificações.',
    'Registrar operações no Firebase com autoria e vínculos. Atualizar consultas enquanto a página está aberta; falhas de gravação não confirmam um salvamento.'
]:
    p(text, 'List Bullet')
p('Cenários disponíveis no simulador', 'Heading 1')
p('Processo dentro dos limites; parâmetro fora da faixa; parada e retomada; refugos e perda de material; meta não atingida; leitura ausente; leitura inválida; limite pendente; dispersão zero; amostra insuficiente; análise pela Engenharia.')
p('Os 11 cenários ficam em memória. Apontamentos e decisões feitos neles não vão para o Firebase. Sair ou recarregar descarta a simulação. Os cálculos usam amostras fictícias e os mesmos serviços do sistema.')
p('Fechamento', 'Heading 1')
p('“Esse é o fluxo que construímos a partir dos materiais. Qual parte se aproxima mais da rotina de vocês e qual precisaríamos ajustar primeiro?”')
p('Limites da apresentação: não temos coleta automática das máquinas, OCR, QR, IA ou fila offline persistente. Importação CSV e comparação ainda não têm fluxo visual completo. Cp/Cpk são estimativas não homologadas. O roteiro descreve as funções da interface verificada em 05/10/2026, não cada função interna do código.')
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
