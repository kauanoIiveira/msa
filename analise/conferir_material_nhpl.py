"""Read original supplied files; never save or modify the workbook."""
from pathlib import Path
import hashlib
import json
import re
import sys
import openpyxl

source = Path(sys.argv[1]).resolve()
target = Path(sys.argv[2]).resolve()
files = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(source.iterdir()) if p.is_file()}
book_path = source / 'T20A03(EN)5 - Capability study senai.xlsx'
formulas = openpyxl.load_workbook(book_path, read_only=True, data_only=False)
cached = openpyxl.load_workbook(book_path, read_only=True, data_only=True)
sheet = formulas['Selo ']
selected = ['A1', 'A2', 'E5', 'I5', 'O5', 'U5', 'E6', 'I6', 'O6', 'U6', 'A7', 'BB9', 'BB10', 'BB11', 'BD9', 'BD10', 'BD11', 'BH9', 'BH10', 'BH11', 'CF9', 'CF10', 'CF11', 'CH9', 'CH10', 'CH11']
cells = {address: {'value_or_formula': sheet[address].value, 'cached': cached['Selo '][address].value} for address in selected}
pattern = re.compile(r'takt|produtiv|\bOEE\b|NHPL|\bMTBF\b|\bMTTR\b|hora|ciclo|tempo|\bT20\b', re.I)
matches = []
for ws in formulas:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and not cell.value.startswith('=') and pattern.search(cell.value):
                matches.append({'sheet': ws.title, 'cell': cell.coordinate, 'text': cell.value})
report = {'read_only': True, 'source_directory': str(source), 'file_sha256': files,
          'sheets': [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in formulas],
          'selected_cells': cells, 'matching_labels': matches,
          'source_unchanged_after_read': all(hashlib.sha256((source/name).read_bytes()).hexdigest() == value for name, value in files.items())}
formulas.close()
cached.close()
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str) + '\n', encoding='utf-8')
print(json.dumps({'report': str(target), 'source_unchanged': report['source_unchanged_after_read'], 'sheets': report['sheets'], 'cells': cells, 'matching_labels': matches}, ensure_ascii=False, default=str))
