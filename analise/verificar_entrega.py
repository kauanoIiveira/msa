from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from decimal import Decimal, localcontext
from pathlib import Path

import openpyxl
import pdfplumber
from pypdf import PdfReader

from gerar_entrega import num


ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "analise" / "evidencias"
DELIVERY = ROOT / "entrega"
PDF = ROOT / "output" / "pdf" / "Preparacao_Entrevista_MSA_06-10-2026.pdf"


def main():
    assert num(10) == "10" and num(0) == "0" and num(0.1) == "0,1"
    sources = json.loads((EVIDENCE / "fontes_sha256.json").read_text(encoding="utf-8"))
    for source in sources:
        path = Path(source["path"])
        assert hashlib.sha256(path.read_bytes()).hexdigest() == source["sha256"], path
    book = openpyxl.load_workbook(Path(sources[0]["path"]), data_only=False)
    sheet = book["Selo "]
    audit = json.loads((EVIDENCE / "auditoria_calculos.json").read_text(encoding="utf-8"))
    summary = audit["summary"]
    assert summary["comparison_count"] == summary["matches"] == 224
    assert not summary["mismatches"]
    assert summary["type_counts"] == {"numero": 562, "texto_decimal": 121, "ausente": 14}
    assert summary["parameters"] == 41 and summary["potential_parameter_slots"] == 697
    assert sheet["BH10"].value == 80 and sheet["BH11"].value == 75
    assert sheet["BP9"].value.strip() == "Retardo de Mesa"
    normal = book["Normality test "]
    cdf_cells = [normal.cell(r, 29) for r in range(10, 1010)]
    assert len(cdf_cells) == 1000
    assert all(c.data_type == "n" and 0 < c.value < 1 for c in cdf_cells)
    assert len({c.value for c in cdf_cells}) > 1
    assert all(normal.cell(r, 2).value is None for r in range(10, 1010))
    assert sheet["G93"].data_type == "s" and sheet["G93"].value == "OK"

    # Decimal is an independent check of the BF arithmetic, not the float/statistics implementation.
    with localcontext() as ctx:
        ctx.prec = 40
        values = [Decimal(str(sheet.cell(r, 58).value).replace(",", ".")) for r in range(17, 34)]
        assert Counter(values) == {Decimal("0.8"): 7, Decimal("0.9"): 10}
        mean = sum(values) / Decimal(len(values))
        variance = sum((x - mean)**2 for x in values) / Decimal(len(values))
        sd = variance.sqrt()
        cp = (Decimal("0.9") - Decimal("0.8")) / (6*sd)
        cpk = min(Decimal("0.9")-mean, mean-Decimal("0.8")) / (3*sd)
        bf = next(p for p in audit["parameters"] if p["column"] == "BF")["normalized_math"]
        for key, exact in [("mean", mean), ("sd_population", sd), ("cp_arithmetic", cp), ("cpk_arithmetic", cpk)]:
            assert abs(Decimal(str(bf[key])) - exact) < Decimal("1e-12"), key
    assert sum([72, 216, 180, 144, 180, 72, 108, 180]) == 1152
    assert Decimal("6.6") * Decimal("0.980665") == Decimal("6.4723890")

    for filename, expected in [("auditoria_41_parametros.csv", 41), ("medicoes_697_campos.csv", 697)]:
        with (DELIVERY / filename).open(encoding="utf-8-sig", newline="") as stream:
            records = list(csv.DictReader(stream, delimiter=";"))
        assert len(records) == expected, filename
        assert all(None not in record for record in records), filename
    md = (DELIVERY / "AUDITORIA_TECNICA_MSA.md").read_text(encoding="utf-8")
    assert "{{" not in md and md.count("#### ") == 41
    for p in audit["parameters"]:
        assert f"#### {p['column']} | {p['name']}" in md

    reader = PdfReader(PDF)
    assert len(reader.pages) == 8
    page_texts = [p.extract_text() for p in reader.pages]
    expected_titles = ["01 / A entrevista", "02 / Uma evidência", "03 / Regras antes", "04 / O contexto", "05 / Produção", "06 / Uso real", "07 / O que entra", "08 / Sair com"]
    for text, title in zip(page_texts, expected_titles):
        assert title in text, title
        assert "\ufffd" not in text and "\u25a0" not in text
        assert len(text) > 1000
    with pdfplumber.open(PDF) as doc:
        outside = []
        for i, page in enumerate(doc.pages):
            outside.extend((i+1, c["text"]) for c in page.chars if c["x0"] < 40 or c["x1"] > page.width-39 or c["top"] < 17 or c["bottom"] > page.height-17)
            for img in page.images:
                assert img["x0"] >= 40 and img["x1"] <= page.width-39, (i+1, img)
        assert not outside, outside[:10]

    result = {
        "fontes_preservadas_por_sha256": len(sources),
        "resumos_numericos_comparados": 224,
        "resumos_coincidentes": 224,
        "BF_verificacao_independente_Decimal": {"media": str(mean), "desvio_populacional": str(sd), "cp_aritmetico": str(cp), "cpk_aritmetico": str(cpk)},
        "csv_parametros": 41,
        "csv_campos": 697,
        "parametros_no_relatorio": 41,
        "paginas_pdf": 8,
        "texto_e_imagens_dentro_das_margens": True,
        "observacao": "Verificacao estrutural/numerica; a inspecao visual das paginas renderizadas e realizada separadamente.",
    }
    (EVIDENCE / "verificacao_entrega.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
