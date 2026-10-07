from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import statistics
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\Kauan\Desktop\MSA\T20A03(EN)5 - Capability study senai.xlsx")
OUT = ROOT / "analise" / "evidencias"
DELIVERY = ROOT / "entrega"


def parse_number(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value), "numero"
    if isinstance(value, str):
        text = value.strip()
        if re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", text):
            return float(text.replace(",", ".")), "texto_decimal"
        return None, "texto_nao_numerico"
    return None, "ausente"


def format_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%d/%m/%Y").strftime("%Y-%m-%d")
        except ValueError:
            return value
    return None


def metrics(values, low, high):
    if not values:
        return {}
    mean = statistics.mean(values)
    sd = statistics.pstdev(values)
    result = {"n": len(values), "min": min(values), "max": max(values),
              "mean": mean, "sd_population": sd,
              "sd_sample": statistics.stdev(values) if len(values) > 1 else None}
    if low is not None and high is not None and sd > 0:
        result["cp_arithmetic"] = (high - low) / (6 * sd)
        result["cpk_arithmetic"] = min(high - mean, mean - low) / (3 * sd)
    else:
        result["cp_arithmetic"] = None
        result["cpk_arithmetic"] = None
    return result


def save_json(name, obj):
    (OUT / name).write_text(json.dumps(obj, indent=2, ensure_ascii=False, default=str), encoding="utf-8")


def save_csv(name, rows):
    with (DELIVERY / name).open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, list(rows[0]), delimiter=";")
        writer.writeheader()
        writer.writerows(rows)


def main():
    DELIVERY.mkdir(parents=True, exist_ok=True)
    book = openpyxl.load_workbook(SOURCE, data_only=False)
    cached = openpyxl.load_workbook(SOURCE, data_only=True)
    sheet = book["Selo "]
    saved = cached["Selo "]
    columns = [c for c in range(6, sheet.max_column + 1) if sheet.cell(9, c).value is not None]
    assert len(columns) == 41
    rows = [r for r in range(17, 68) if any(sheet.cell(r, c).value is not None for c in columns)]
    assert rows == list(range(17, 34))
    parameters, long_records, comparisons, flat = [], [], [], []
    for column in columns:
        col = get_column_letter(column)
        low, _ = parse_number(sheet.cell(10, column).value)
        high, _ = parse_number(sheet.cell(11, column).value)
        name = " ".join(str(sheet.cell(9, column).value).split())
        original_name = sheet.cell(9, column).value
        numeric, normalized, records = [], [], []
        for row in rows:
            cell = sheet.cell(row, column)
            number, status = parse_number(cell.value)
            if status == "numero":
                numeric.append(number)
            if number is not None:
                normalized.append(number)
            relation = "nao_avaliado"
            if number is not None and low is not None and high is not None and low < high:
                relation = "abaixo" if number < low else "acima" if number > high else "dentro"
            record = {"celula": cell.coordinate, "data": format_date(sheet.cell(row, 1).value),
                      "coluna": col, "parametro": name, "unidade": sheet.cell(12, column).value,
                      "original": cell.value, "tipo_excel": cell.data_type,
                      "formato_excel": cell.number_format, "valor_interpretado": number,
                      "tratamento": status, "relacao_faixa_do_arquivo": relation}
            records.append(record)
            long_records.append(record)
        native = metrics(numeric, low, high)
        all_values = metrics(normalized, low, high)
        cached_metrics = {"min": saved.cell(69, column).value, "max": saved.cell(70, column).value,
                          "mean": saved.cell(71, column).value, "sd_population": saved.cell(72, column).value,
                          "cp_arithmetic": saved.cell(73, column).value,
                          "cpk_arithmetic": saved.cell(74, column).value,
                          "status": saved.cell(75, column).value}
        issues = []
        types = Counter(r["tratamento"] for r in records)
        if types["texto_decimal"]:
            issues.append(f"{types['texto_decimal']} valores como texto")
        if low is not None and high is not None and low > high:
            issues.append("minimo maior que maximo")
        elif low is not None and high is not None and low == high:
            issues.append("limites iguais")
        elif low is None or high is None:
            issues.append("especificacao unilateral ou incompleta")
        if sheet.cell(17, column).value is None and normalized:
            issues.append("primeira linha vazia suprime resumo")
        if native and native["sd_population"] == 0:
            issues.append("dispersao matematica nula no subconjunto numerico")
        if all_values and all_values["sd_population"] == 0:
            issues.append("dispersao matematica nula em todos os registros interpretados")
        if any(isinstance(v, str) and v.startswith("#") for v in cached_metrics.values()):
            issues.append("erro de formula salvo")
        for metric, expected in native.items():
            if metric not in cached_metrics or expected is None:
                continue
            observed = cached_metrics[metric]
            if isinstance(observed, (int, float)):
                ok = math.isclose(expected, observed, rel_tol=1e-9, abs_tol=1e-9)
                comparisons.append({"cell": f"{col}{ {'min':69,'max':70,'mean':71,'sd_population':72,'cp_arithmetic':73,'cpk_arithmetic':74}[metric] }",
                                    "metric": metric, "calculated": expected, "cached": observed,
                                    "match": ok, "difference": expected - observed})
        hist = [saved.cell(r, column).value for r in range(98, 110)]
        hist_total = sum(v for v in hist if isinstance(v, (int, float)))
        param = {"column": col, "source_number": sheet.cell(8, column).value,
                 "name": name, "original_name": original_name, "unit": sheet.cell(12, column).value,
                 "low": low, "high": high, "types": dict(types), "cached": cached_metrics,
                 "native_math": native, "normalized_math": all_values,
                 "counts_against_file_limits": dict(Counter(r["relacao_faixa_do_arquivo"] for r in records)),
                 "histogram_saved_sum": hist_total, "issues": issues,
                 "formulas": {f"{col}{r}": sheet.cell(r, column).value for r in range(69, 76)},
                 "records": records}
        parameters.append(param)
        flat.append({"coluna": col, "parametro": name, "unidade": param["unit"],
                     "minimo_arquivo": low, "maximo_arquivo": high,
                     "numeros_originais": types["numero"], "textos_decimais": types["texto_decimal"],
                     "ausentes": types["ausente"], "n_interpretado": all_values.get("n"),
                     "media_salva": cached_metrics["mean"], "media_interpretada": all_values.get("mean"),
                     "desvio_pop_salvo": cached_metrics["sd_population"],
                     "desvio_pop_interpretado": all_values.get("sd_population"),
                     "desvio_amostral_interpretado": all_values.get("sd_sample"),
                     "cp_salvo": cached_metrics["cp_arithmetic"],
                     "cp_aritmetico_interpretado": all_values.get("cp_arithmetic"),
                     "cpk_salvo": cached_metrics["cpk_arithmetic"],
                     "cpk_aritmetico_interpretado": all_values.get("cpk_arithmetic"),
                     "status_salvo": cached_metrics["status"],
                     "histograma_contagem_salva": hist_total, "observacoes": " / ".join(issues)})

    normal = book["Normality test "]
    normal_saved = cached["Normality test "]
    normal_audit = {
        "input_nonempty": sum(normal.cell(r, 2).value is not None for r in range(10, 1010)),
        "n_cached": normal_saved["AE6"].value,
        "k5": {"value": normal["K5"].value, "type": normal["K5"].data_type},
        "h5_formula": normal["H5"].value, "h7_formula": normal["H7"].value,
        "g93": {"value": sheet["G93"].value, "type": sheet["G93"].data_type},
        "reverse_cdf_numeric_constants": sum(normal.cell(r, 29).data_type == "n" and normal.cell(r, 29).value is not None for r in range(10, 1010)),
        "reverse_cdf_examples": {"AC10": normal["AC10"].value, "AC1009": normal["AC1009"].value},
        "count_formula": normal["AE6"].value,
        "input_sort_formula_present": any(isinstance(c.value, str) and any(x in c.value.upper() for x in ["SORT(", "SMALL(", "LARGE("])
                                           for row in normal for c in row),
        "normal_sheet_references_in_selo": [c.coordinate for row in sheet for c in row if c.data_type == "f" and "Normality" in c.value],
    }
    errors = [{"sheet": s.title, "cell": c.coordinate, "error": cached[s.title][c.coordinate].value,
               "formula": c.value} for s in book for row in s for c in row
              if cached[s.title][c.coordinate].data_type == "e"]
    text_by_row = Counter(int(re.search(r"\d+", r["celula"])[0]) for r in long_records if r["tratamento"] == "texto_decimal")
    summary = {"source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
               "parameters": len(parameters), "collection_rows": rows, "collection_count": len(rows),
               "potential_parameter_slots": len(rows) * len(columns),
               "type_counts": dict(Counter(r["tratamento"] for r in long_records)),
               "text_counts_by_row": dict(text_by_row),
               "status_counts": dict(Counter(p["cached"]["status"] for p in parameters)),
               "comparison_count": len(comparisons),
               "matches": sum(c["match"] for c in comparisons),
               "mismatches": [c for c in comparisons if not c["match"]],
               "normality": normal_audit,
               "metadata_nonempty": {get_column_letter(c): sum(sheet.cell(r, c).value is not None for r in rows) for c in range(2, 6)},
               "error_count": len(errors),
               "formula_count": sum(c.data_type == "f" for s in book for row in s for c in row)}
    save_json("auditoria_calculos.json", {"summary": summary, "parameters": parameters,
                                          "comparisons": comparisons, "errors": errors})
    save_csv("auditoria_41_parametros.csv", flat)
    save_csv("medicoes_697_campos.csv", long_records)
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    for p in parameters:
        print(p["column"], p["name"], p["low"], p["high"], p["types"],
              "saved=", p["cached"]["cp_arithmetic"], p["cached"]["cpk_arithmetic"],
              "all=", p["normalized_math"].get("cp_arithmetic"), p["normalized_math"].get("cpk_arithmetic"),
              "hist=", p["histogram_saved_sum"])


if __name__ == "__main__":
    main()
