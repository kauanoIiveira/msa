from __future__ import annotations

import hashlib
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\Kauan\Desktop\MSA")
OUT = ROOT / "analise" / "evidencias"
PDFS = [
    Path(r"C:\Users\Kauan\Downloads\Analise_Materiais_MSA_2026.pdf"),
    Path(r"C:\Users\Kauan\Downloads\Dossie_MSA_SENAI_2026.pdf"),
]


def serial(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def write_json(name, value):
    (OUT / name).write_text(
        json.dumps(value, ensure_ascii=False, indent=2, default=serial),
        encoding="utf-8",
    )


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    paths = sorted(SOURCE.iterdir()) + PDFS
    manifest = [
        {"path": str(p), "bytes": p.stat().st_size,
         "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
        for p in paths if p.is_file()
    ]
    write_json("fontes_sha256.json", manifest)

    xlsx = next(SOURCE.glob("*.xlsx"))
    book = openpyxl.load_workbook(xlsx, data_only=False)
    cached = openpyxl.load_workbook(xlsx, data_only=True)
    overview = {"source": str(xlsx), "calculation": str(book.calculation),
                "defined_names": {k: str(v) for k, v in book.defined_names.items()},
                "sheets": []}
    for sheet in book:
        cached_sheet = cached[sheet.title]
        cells = []
        for row in sheet:
            for cell in row:
                if cell.value is None:
                    continue
                cells.append({"address": cell.coordinate, "value": cell.value,
                              "type": cell.data_type, "format": cell.number_format,
                              "cached": cached_sheet[cell.coordinate].value,
                              "cached_type": cached_sheet[cell.coordinate].data_type,
                              "comment": cell.comment.text if cell.comment else None})
        safe_name = sheet.title.strip().replace(" ", "_")
        write_json(f"celulas_{safe_name}.json", cells)
        summary = {"name": sheet.title, "state": sheet.sheet_state,
                   "dimensions": sheet.calculate_dimension(), "cells": len(cells),
                   "formula_count": sum(c["type"] == "f" for c in cells),
                   "cache_errors": dict(Counter(c["cached"] for c in cells if c["cached_type"] == "e")),
                   "merged": [str(r) for r in sheet.merged_cells.ranges],
                   "hidden_rows": [r for r, d in sheet.row_dimensions.items() if d.hidden],
                   "hidden_cols": [c for c, d in sheet.column_dimensions.items() if d.hidden],
                   "validation": str(sheet.data_validations),
                   "tables": list(sheet.tables), "charts": len(sheet._charts)}
        overview["sheets"].append(summary)
        print(json.dumps(summary, ensure_ascii=False))
        if sheet.title.strip() == "Selo":
            for row in list(range(1, 17)) + list(range(68, min(sheet.max_row, 100) + 1)):
                values = [f"{c.coordinate}={c.value!r}" for c in sheet[row] if c.value is not None]
                print(" | ".join(values[:8]) + (f" | ... {len(values)} cells" if len(values) > 8 else ""))
        else:
            for row in range(1, 11):
                values = [f"{c.coordinate}={c.value!r} [cache={cached_sheet[c.coordinate].value!r}]"
                          for c in sheet[row] if c.value is not None]
                print(" | ".join(values[:12]))
    write_json("estrutura_planilha.json", overview)

    for path in PDFS:
        reader = PdfReader(path)
        text = "\n\n".join(f"--- PAGINA {i + 1} ---\n{page.extract_text()}"
                              for i, page in enumerate(reader.pages))
        (OUT / (path.stem + ".txt")).write_text(text, encoding="utf-8")
        print(json.dumps({"pdf": path.name, "pages": len(reader.pages),
                          "characters": len(text), "metadata": str(reader.metadata)}, ensure_ascii=False))

    photos = sorted(SOURCE.glob("*.jpeg"))
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 21)
    for batch in range(0, len(photos), 3):
        canvas = Image.new("RGB", (1500, 1030), "#f1f3f5")
        draw = ImageDraw.Draw(canvas)
        for offset, path in enumerate(photos[batch:batch + 3]):
            img = Image.open(path)
            print(json.dumps({"photo": path.name, "size": img.size}))
            img.thumbnail((480, 950))
            x = offset * 500 + (500 - img.width) // 2
            canvas.paste(img, (x, 60))
            draw.text((offset * 500 + 15, 15), f"Foto {batch + offset + 1} - {path.stem[-8:]}",
                      fill="#17212b", font=font)
        canvas.save(OUT / f"fotos_{batch + 1}_{min(batch + 3, len(photos))}.jpg")

    crops = [
        (photos[1], (20, 340, 875, 775), "pitch_board_detalhe.png"),
        (photos[7], (375, 180, 600, 325), "aquecimento_detalhe.png"),
    ]
    for path, box, name in crops:
        crop = Image.open(path).crop(box)
        crop.resize((crop.width * 3, crop.height * 3)).save(OUT / name)


if __name__ == "__main__":
    main()
