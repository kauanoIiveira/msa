from pathlib import Path
import hashlib, json, re, shutil, statistics, subprocess
from datetime import datetime, timezone
from collections import Counter
from openpyxl import load_workbook

ROOT = Path(r'C:\Users\Kauan\Documents\ChatGPT\Desafio de Ideias')
SOURCE = Path(r'C:\Users\Kauan\Desktop\Grupo Amarelo - MSA Brasil')
DESKTOP = Path(r'C:\Users\Kauan\Desktop\msa-master')
OUT = ROOT / 'docs/continuidade/2026-10-06'
EVID = OUT / 'evidencias'
MEMORY_NOTE = Path(r'C:\Users\Kauan\.codex\memories\extensions\ad_hoc\notes\2026-10-06T18-10-00-msa-conferencia-fontes.md')

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def readjson(name):
    return json.loads((EVID/name).read_text(encoding='utf-8'))

def savejson(name, obj):
    (EVID/name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding='utf-8')

def link(label, path):
    return f'[{label}](<{str(path).replace(chr(92), chr(47))}>)'

# Preserve exact original byte slices, including spacing and blank lines.
raw = (OUT/'PROMPT_ORIGINAL_INTEGRAL.txt').read_bytes()
start = raw.index('Sou Coordenadora Técnica'.encode('utf-8'))
end = raw.index('Perguntas feitas:'.encode('utf-8'), start)
(OUT/'SOBRE_FABIANA_ORIGINAL.txt').write_bytes(raw[start:end])
start = end
end = raw.index('Analise TODO o conteúdo do Drive'.encode('utf-8'), start)
(OUT/'PERGUNTAS_ORIGINAIS.txt').write_bytes(raw[start:end])

# Preserve previous source documents; the Desktop CONTINUE body is the original snapshot.
prior_dir = EVID/'documentacao_anterior'
prior_dir.mkdir(exist_ok=True)
source_mds = ['AUDITORIA_TECNICA_MSA.md', 'ROTEIRO_ENTREVISTA_MSA.md']
for name in source_mds:
    shutil.copyfile(SOURCE/'Rascunhos Persona'/name, prior_dir/name)
doc_paths = ['README.md', 'CONTINUE_AQUI.md', 'docs/CONTRATOS_FUNCIONAIS.md',
             'docs/ENTREGA_FUNCIONAL.md', 'docs/PARAMETROS_MSA.md',
             'docs/superpowers/specs/2026-10-05-msa-firebase-design.md',
             'docs/superpowers/plans/2026-10-05-msa-firebase-implementation.md']
for rel in doc_paths:
    shutil.copyfile(DESKTOP/rel, prior_dir/Path(rel).name)
script_dir = EVID/'scripts'
script_dir.mkdir(exist_ok=True)
for name in ['consolidar_fontes_msa.py','revisar_continuidade_msa.py','finalizar_continuidade_msa.py']:
    shutil.copyfile(ROOT/'tmp'/name, script_dir/name)

# Independent read-only recalculation of the BF example and field classification.
book_path = SOURCE/'MSA - Material Fornecido/T20A03(EN)5 - Capability study senai.xlsx'
book = load_workbook(book_path, data_only=False, read_only=False)
sheet = book['Selo ']
cols = list(range(6,87,2))
time_columns = [sheet.cell(9,c).coordinate for c in cols if sheet.cell(12,c).value == 'seg']
values = [float(str(sheet.cell(r,58).value).replace(',','.')) for r in range(17,34)]
native = [sheet.cell(r,58).value for r in range(17,34) if isinstance(sheet.cell(r,58).value,(int,float))]
mean = statistics.mean(values)
sigma = statistics.pstdev(values)
lo, hi = sheet['BF10'].value, sheet['BF11'].value
example = dict(source=str(book_path),sheet=sheet.title,range='BF17:BF33',n=len(values),
    native_count=len(native),text_count=len(values)-len(native),frequency=dict(Counter(map(str,values))),
    mean=mean,population_sigma=sigma,cp_arithmetic=(hi-lo)/(6*sigma),
    cpk_arithmetic=min(hi-mean,mean-lo)/(3*sigma),homologated=False,
    note='Recálculo independente de exemplo, não homologação industrial; entradas originais não alteradas.',
    fields_with_unit_seg=time_columns,count_fields_with_unit_seg=len(time_columns))
assert len(values)==17 and len(native)==7 and len(time_columns)==15
savejson('recalculo_exemplo_planilha.json',example)
book.close()

inventory = readjson('inventario_local.json')
drive = readjson('inventario_drive.json')
drive_compare = readjson('conferencia_drive_local.json')
compare = readjson('comparacao_checkouts.json')
docx = {i['path']:i['extraction'] for i in readjson('indice_docx.json')}
assert len(inventory)==30 and len(drive)==31 and len(drive_compare)==27
assert all(x['identical_local_files'] for x in drive_compare)

def status(item):
    ext=item['type']; rel=item['relative']; name=Path(item['path']).name
    if ext=='.m4a':
        return ('PENDENTE — apenas arquivo/hash/cabeçalho', 'evidencias/audio_pendencia.json', 'Entrevista: ouvir e transcrever; cópias idênticas')
    if rel.startswith('Desafio\\'):
        return ('Lido visualmente e transcrito', 'ENUNCIADO_TRANSCRITO.md', 'Problema e 13 mínimos/16 opcionais; final termina em dois-pontos')
    if ext=='.xlsx':
        return ('Todas as abas/células preenchidas/fórmulas extraídas e revisadas', 'evidencias/indice_planilha.json', '41 referências/17 linhas/697 campos; homologação e semântica pendentes')
    if ext=='.pdf':
        pages={'Analise_Materiais_MSA_2026.pdf':22,'Dossie_MSA_SENAI_2026.pdf':17,'Preparacao_Entrevista_MSA_06-10-2026.pdf':8}[name]
        return (f'Texto integral lido + revisão visual de {pages} páginas em pranchas', f'evidencias/{Path(name).stem}.txt', 'Análise/rascunho da equipe; não depoimento validado')
    if ext=='.docx':
        return ('Conteúdo textual integral OOXML lido; sem mídia interna', 'evidencias/'+docx[item['path']], 'Roteiro/perguntas/referência/proposta; layout DOCX não re-renderizado nesta etapa')
    if ext=='.csv':
        n=41 if '41' in name else 697
        return (f'Todas as {n} linhas de dados lidas e cópia conferida', 'evidencias/indice_csv.json', 'Extração anterior rastreável; idêntica à entrega local')
    if ext=='.md':
        return ('Texto integral lido e preservado', 'evidencias/documentacao_anterior/'+name, 'Auditoria/roteiro anteriores; separar parecer de resposta da empresa')
    if rel.startswith('Perguntas para o Persona\\'):
        return ('Lido visualmente e transcrito', 'PERGUNTAS_IMAGENS_TRANSCRITAS.md', 'Perguntas anteriores sem respostas confirmadas')
    if rel.startswith('Matriz de Prioridade\\'):
        return ('Lido visualmente', 'CONSOLIDACAO.md, seção5', 'Priorização da equipe, não rubrica oficial')
    return ('Lido visualmente, individualmente', 'CONSOLIDACAO.md, seção4', 'IHM/processo/instrumentos/papel; incertezas documentadas')

lines=['# Inventário de fontes e cobertura — 06/10/2026','',
    'A análise documental foi concluída com as fontes acessíveis, com uma fonte lógica relevante pendente: a entrevista em áudio. Os estados abaixo distinguem leitura de conteúdo, conferência de arquivo e exclusão autorizada. Os hashes estão nos inventários JSON e foram novamente conferidos ao finalizar.','',
    '## Caminhos e links fornecidos, preservados','',
    '- Drive: https://drive.google.com/drive/folders/1wj03b32ZzVM-7Tj53Dh2VYi_4KI7X9LL?usp=sharing',
    '- Material local: `C:\\Users\\Kauan\\Desktop\\Grupo Amarelo - MSA Brasil`',
    '- Projeto local já existente: `C:\\Users\\Kauan\\Desktop\\msa-master`',
    '- GitHub: https://github.com/kauanoIiveira/msa',
    '- Rascunhos de persona: `C:\\Users\\Kauan\\Desktop\\Grupo Amarelo - MSA Brasil\\Rascunhos Persona`','',
    'O pedido chama a fonte de “Material Fornecido pela empresa”; no disco/Drive o nome encontrado é “MSA - Material Fornecido”. Todas as fontes acima foram acessíveis. O checkout ativo é `C:\\Users\\Kauan\\Documents\\ChatGPT\\Desafio de Ideias`. A gravação está acessível como arquivo, mas seu conteúdo não pôde ser processado.','',
    '## Arquivos locais — 30','',
    '| ID | Caminho exato/tipo | Estado da leitura | Extração/referência | Relação e pendência |',
    '|---|---|---|---|---|']
local_ids={}
for i,item in enumerate(inventory,1):
    ident=f'L{i:02d}'; local_ids[item['path']]=ident
    s,e,r=status(item)
    lines.append(f'| {ident} | {link(item["relative"],Path(item["path"]))}<br>{item["type"]} | {s} | {e} | {r} |')
lines+=['','## Drive — 31 itens de arquivo','',
    'A contagem não inclui pastas. Foi percorrida a árvore de pastas e materializado o conteúdo dos 27 arquivos binários; todos têm SHA-256 idêntico a uma fonte local. Três documentos nativos foram lidos por texto. O ZIP do MVP foi apenas inventariado, sem download ou análise, conforme a única exclusão solicitada. Títulos com espaço inicial e caracteres Unicode foram preservados no inventário JSON; a tabela usa o mesmo título.','',
    '| ID | Pasta e título original/link | Tipo | Estado / vínculo local |',
    '|---|---|---|---|']
drive_byid={item['id']:item for item in drive_compare}
native_index={
    'Lista de Materiais e ganho de produção':('evidencias/drive_documento_1.txt','Lista de Materiais e ganho de produção.docx'),
    'O que deve conter no Dashboard':('evidencias/drive_documento_2.txt','Wireframes\\O que deve conter no Dashboard.docx'),
    'Perguntas para o Persona':('evidencias/drive_documento_3.txt','Perguntas para o Persona\\Perguntas para o Persona.docx')}
zip_count=0
for i,item in enumerate(drive,1):
    ident=f'D{i:02d}'
    if item['title']=='msa-master.zip':
        state='EXCLUÍDO por instrução explícita; protótipo analisado via checkout/Desktop/GitHub';zip_count+=1
    elif item['mime_type']=='application/vnd.google-apps.document':
        extraction,localrel=native_index[item['title']]
        state=f'Texto nativo integral lido: {extraction}; exportação local {localrel}'
    else:
        matches=drive_byid[item['id']]['identical_local_files']
        ids=', '.join(local_ids[p] for p in matches)
        state=f'Binário conferido por SHA-256 com {ids}; mesmo estado de leitura da fonte local'
    lines.append(f'| {ident} | {item["folder"]}: [{item["title"]}]({item["url"]}) | {item["mime_type"]} | {state} |')
assert zip_count==1
lines+=['','## Documentação, extrações anteriores e código','',
    '| Fonte | Estado e uso | Referência preservada |', '|---|---|---|']
for rel in doc_paths:
    lines.append(f'| {link(rel,ROOT/rel)} | Texto lido e cruzado com código/fontes; etapa histórica e divergências marcadas | evidencias/documentacao_anterior/{Path(rel).name} |')
for name in ['Roteiro_Entrevista_MSA_Equipe_Amarela.docx','Roteiro_Apresentacao_Funcoes_MSA.docx']:
    lines.append(f'| {link("docs/"+name,ROOT/"docs"/name)} | Conteúdo textual lido; idêntico à cópia de material; original preservado | evidencias/indice_docx.json |')
lines+=['| `analise/` e `entrega/` | Extração anterior recuperada; 12 hashes de fontes conferidos; CSVs conferidos integralmente | evidencias/conferencia_extracoes_anteriores.json e indice_csv.json |',
    '| `app/`, serviços/domínio/repositório/catálogo/UI, regras locais e testes | Código necessário para telas, contexto, coleta, limites, indicadores, Engenharia, histórico, CSV, simulação e permissões revisado; sem escrita no aplicativo | CONSOLIDACAO.md, seções6–8; manifesto de preservação |',
    '| GitHub e cópia Desktop | Referências Git e conteúdo comparados; 117 arquivos correspondentes inicialmente | evidencias/comparacao_checkouts.json |','',
    '## Lacunas específicas de leitura e acesso','',
    '- **Áudio:** duas cópias idênticas, cerca de5min42s pelo cabeçalho; não ouvido, sem transcrição, falantes ou marcações de conteúdo confirmados. Limitação de ferramentas nesta sessão; continuar a análise das demais fontes foi autorizado no pedido.',
    '- **DOCX:** leitura textual integral dos cinco documentos únicos; não houve renderização nova do layout. Nenhum deles contém mídia interna. Não foram editados.',
    '- **Fotos:** manuscrito do planejamento/produto no Pitch Board e unidade parcial do manômetro não permitem precisão total; colunas de temperatura sem legenda confirmada. Estas incertezas estão na seção4 da consolidação.',
    '- **Folha:** termina com dois-pontos, sem continuação visível; não traz rubrica de pesos, prazo oficial ou duração de apresentação.',
    '- **Nuvem/equipamento:** não consultados dados atuais/regras implantadas do Firebase, CLP, rede ou disponibilidade de sensores. Isso é limite de validação operacional, não arquivo omitido da leitura.',
    '- **Histórico:** documentação e memória recuperadas não equivalem à íntegra de todas as conversas anteriores. Decisões comprovadas foram conferidas no código; demais lacunas permanecem explícitas.','',
    '## Critério de integridade','',
    'O prompt integral e os trechos literais foram conferidos byte a byte. Nenhum arquivo original da pasta Grupo Amarelo foi alterado. Os 27 arquivos do Drive coincidem com o inventário local. As cópias anteriores de Markdown preservam texto/perguntas. O aplicativo e os arquivos não rastreados já existentes foram preservados. A verificação final está em `evidencias/registro_verificacao.json`. Esta consolidação não deve ser descrita como transcrição concluída da entrevista.','']
(OUT/'FONTES_E_COBERTURA.md').write_text('\n'.join(lines),encoding='utf-8')

# Final preservation and content checks. Tests were executed earlier this same session.
original_prompt = Path(r'C:\Users\Kauan\.codex\attachments\989777ab-36cb-43d2-9933-7404448bbfb0\Texto colado.txt')
assert sha(original_prompt)==sha(OUT/'PROMPT_ORIGINAL_INTEGRAL.txt')
assert (OUT/'SOBRE_FABIANA_ORIGINAL.txt').read_bytes() in raw
assert (OUT/'PERGUNTAS_ORIGINAIS.txt').read_bytes() in raw
source_integrity=[{'path':x['path'],'same_sha256':sha(Path(x['path']))==x['sha256']} for x in inventory]
assert all(x['same_sha256'] for x in source_integrity)
appfiles=subprocess.check_output(['git','ls-files','app','index.html','firebase'],cwd=ROOT,text=True).splitlines()
code_integrity=[]
for rel in appfiles:
    before=(DESKTOP/rel).read_bytes().replace(b'\r\n',b'\n')
    after=(ROOT/rel).read_bytes().replace(b'\r\n',b'\n')
    assert before==after, rel
    code_integrity.append({'path':rel,'sha256_current':sha(ROOT/rel),'same_content_as_desktop':True})
savejson('manifesto_preservacao_aplicativo.json',code_integrity)
old_continue=(DESKTOP/'CONTINUE_AQUI.md').read_text(encoding='utf-8')
new_continue=(ROOT/'CONTINUE_AQUI.md').read_text(encoding='utf-8')
prefix='# Continuidade do MSA\n\n'
after_marker=new_continue.index('## Abrir em outro computador')
assert prefix+new_continue[after_marker:]==old_continue
report=(OUT/'CONSOLIDACAO.md').read_text(encoding='utf-8')
assert len(re.findall(r'^\| M\d{2} \|',report,re.M))==13
assert len(re.findall(r'^\| O\d{2} \|',report,re.M))==16
questions=(OUT/'PERGUNTAS_E_PENDENCIAS.md').read_text(encoding='utf-8')
assert len(re.findall(r'^\| (Quais|Qual|Como|Em quais|O que|A empresa)',questions,re.M))==12
assert MEMORY_NOTE.is_file()
for file in OUT.rglob('*'):
    if file.is_file() and file.suffix in ['.md','.txt']:
        text=file.read_text(encoding='utf-8')
        assert text.strip(),file
        assert '\ufffd' not in text,file
git_status=subprocess.check_output(['git','status','--short'],cwd=ROOT,text=True).splitlines()
assert git_status==[' M CONTINUE_AQUI.md','?? docs/Roteiro_Apresentacao_Funcoes_MSA.docx','?? scripts/create-presentation-doc.py'],git_status
verification={
    'created_utc':datetime.now(timezone.utc).isoformat(),
    'checkout':str(ROOT),'branch':'master','head':'a268ddcd841ed7cb71f9110752498f0902e16222',
    'prompt_original_sha256':sha(original_prompt),'prompt_copy_identical':True,
    'literal_profile_and_questions_preserved':True,
    'all_30_local_originals_unchanged':True,'source_integrity':source_integrity,
    'drive_files':31,'drive_binary_identical_to_local':27,'drive_native_documents_read':3,'excluded_zip_count':1,
    'local_files':30,'audio_content_pending':True,'logical_audio_sources':1,
    'desktop_comparison_initial':{'files':117,'identical_bytes':48,'line_endings_only':69,'content_differences':0},
    'requirements_rows':13,'optional_rows':16,'original_question_groups':12,
    'app_and_firebase_code_unchanged':True,'continue_original_body_preserved':True,
    'existing_untracked_files_preserved':True,'git_status':git_status,
    'tests_observed_this_session':{'command':'npm test','tests':83,'pass':83,'fail':0,'skip':0,'exit_code':0},
    'static_observed_this_session':{'command':'npm run verify:static -- --deploy','ok':True,'files':45,'errors':[],'exit_code':0},
    'not_repeated_this_session':['browser e2e','Firebase emulator','live Firebase reads/writes','industrial validation'],
    'memory_update_note':str(MEMORY_NOTE),'memory_note_sha256':sha(MEMORY_NOTE),
    'local_git_exclusion':'docs/continuidade and temporary source copies via .git/info/exclude',
    'commit_push_deploy_or_messages_sent':False,
    'outputs':[str(f.relative_to(OUT)) for f in sorted(OUT.rglob('*')) if f.is_file()]}
savejson('registro_verificacao.json',verification)
print(json.dumps({'ok':True,'documents_directory':str(OUT),'local_originals_unchanged':30,
    'application_files_preserved':len(code_integrity),'requirements':13,'optionals':16,
    'questions':12,'memory_note':str(MEMORY_NOTE),'git_status':git_status},ensure_ascii=False,indent=2))
